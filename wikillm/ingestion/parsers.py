"""Multi-format document parsing with robust error handling.

Supported formats and the library used for each:

================  ==============================  ============================
Extension         Library                         Notes
================  ==============================  ============================
``.txt``          ``charset-normalizer``          Encoding auto-detection.
``.md``           ``markdown-it-py``              Markup stripped to plain text.
``.pdf``          ``pypdf``                       Per-page text extraction.
``.docx``         ``python-docx``                 Paragraphs + table cells.
``.doc``          external (``antiword``)         Legacy, best-effort fallback.
================  ==============================  ============================

The two entry points are :func:`parse_bytes` (for uploads / in-memory data) and
:func:`parse_path` (for filesystem ingestion). Both return a fully populated
:class:`~wikillm.models.ParsedDocument`. Parsing problems raise
:class:`ParseError`; unknown formats raise :class:`UnsupportedFormatError`.
"""

from __future__ import annotations

import io
import mimetypes
import os
import shutil
import subprocess
import uuid
from pathlib import Path
from typing import Iterator, Optional

from wikillm.logging_setup import get_logger
from wikillm.models import ParsedDocument

logger = get_logger(__name__)

# Logical content type keyed by lower-case file extension.
SUPPORTED_EXTENSIONS: dict[str, str] = {
    ".txt": "txt",
    ".text": "txt",
    ".md": "md",
    ".markdown": "md",
    ".pdf": "pdf",
    ".docx": "docx",
    ".doc": "doc",
}

# MIME type -> logical content type, used as a secondary signal.
_MIME_TO_TYPE: dict[str, str] = {
    "text/plain": "txt",
    "text/markdown": "md",
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/msword": "doc",
}


class ParseError(Exception):
    """Raised when a supported format cannot be parsed (corrupt/unreadable)."""


class UnsupportedFormatError(Exception):
    """Raised when a file's format is not supported by the ingestion pipeline."""


def detect_type(filename: str, mime_type: Optional[str] = None) -> str:
    """Determine the logical content type for a filename.

    The file extension is the primary signal; the MIME type (if provided or
    guessable) is used as a fallback.

    Args:
        filename: The file name (or path) whose type should be detected.
        mime_type: Optional MIME type hint (e.g. from an HTTP upload).

    Returns:
        A logical content type: ``"txt"``, ``"md"``, ``"pdf"``, ``"docx"`` or
        ``"doc"``.

    Raises:
        UnsupportedFormatError: If neither the extension nor the MIME type maps
            to a supported format.
    """
    ext = Path(filename).suffix.lower()
    if ext in SUPPORTED_EXTENSIONS:
        return SUPPORTED_EXTENSIONS[ext]

    guessed = mime_type or mimetypes.guess_type(filename)[0]
    if guessed and guessed in _MIME_TO_TYPE:
        return _MIME_TO_TYPE[guessed]

    raise UnsupportedFormatError(
        f"Unsupported file format for '{filename}'. "
        f"Supported extensions: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
    )


def parse_bytes(
    data: bytes,
    filename: str,
    mime_type: Optional[str] = None,
    metadata: Optional[dict] = None,
) -> ParsedDocument:
    """Parse raw bytes into a :class:`~wikillm.models.ParsedDocument`.

    Args:
        data: The raw file content.
        filename: Original file name (drives type detection and the title).
        mime_type: Optional MIME hint (e.g. from the upload's ``content-type``).
        metadata: Optional metadata to attach to the resulting document.

    Returns:
        The parsed document with extracted plain text.

    Raises:
        UnsupportedFormatError: If the format is not supported.
        ParseError: If extraction fails (e.g. corrupt PDF, missing converter).
    """
    content_type = detect_type(filename, mime_type)
    text = _extract_text(data, content_type, filename)

    if not text.strip():
        raise ParseError(
            f"No extractable text found in '{filename}'. "
            "The file may be empty, image-only, or a scanned document."
        )

    return ParsedDocument(
        id=uuid.uuid4().hex,
        title=Path(filename).name,
        source=filename,
        content_type=content_type,
        text=text,
        metadata=metadata or {},
    )


def parse_path(path: str | os.PathLike[str], metadata: Optional[dict] = None) -> ParsedDocument:
    """Parse a single file from the filesystem.

    Args:
        path: Path to an existing file.
        metadata: Optional metadata to attach.

    Returns:
        The parsed document.

    Raises:
        FileNotFoundError: If the path does not exist or is not a file.
        UnsupportedFormatError: If the file format is not supported.
        ParseError: If extraction fails.
    """
    file_path = Path(path)
    if not file_path.is_file():
        raise FileNotFoundError(f"Not a file: {file_path}")

    data = file_path.read_bytes()
    document = parse_bytes(data, str(file_path), metadata=metadata)
    # For filesystem sources, prefer the absolute path as the source.
    document.source = str(file_path.resolve())
    return document


def iter_supported_files(root: str | os.PathLike[str], recursive: bool = True) -> Iterator[Path]:
    """Yield supported files under a directory (or the file itself).

    Args:
        root: A directory or single file path.
        recursive: When ``root`` is a directory, whether to walk subdirectories.

    Yields:
        Paths to files whose extension is supported.

    Raises:
        FileNotFoundError: If ``root`` does not exist.
    """
    root_path = Path(root)
    if root_path.is_file():
        if root_path.suffix.lower() in SUPPORTED_EXTENSIONS:
            yield root_path
        return
    if not root_path.is_dir():
        raise FileNotFoundError(f"Path does not exist: {root_path}")

    iterator = root_path.rglob("*") if recursive else root_path.glob("*")
    for candidate in sorted(iterator):
        if candidate.is_file() and candidate.suffix.lower() in SUPPORTED_EXTENSIONS:
            yield candidate


# ---------------------------------------------------------------------------
# Per-format extractors
# ---------------------------------------------------------------------------
def _extract_text(data: bytes, content_type: str, filename: str) -> str:
    """Dispatch to the correct extractor for a logical content type.

    Args:
        data: Raw file bytes.
        content_type: Logical type from :func:`detect_type`.
        filename: Original filename (used in error messages).

    Returns:
        The extracted plain text.

    Raises:
        ParseError: If extraction fails.
        UnsupportedFormatError: If the content type has no extractor.
    """
    if content_type == "txt":
        return _extract_text_plain(data)
    if content_type == "md":
        return _extract_text_markdown(data)
    if content_type == "pdf":
        return _extract_text_pdf(data, filename)
    if content_type == "docx":
        return _extract_text_docx(data, filename)
    if content_type == "doc":
        return _extract_text_doc(data, filename)
    raise UnsupportedFormatError(f"No extractor for content type '{content_type}'.")


def _decode_text(data: bytes) -> str:
    """Decode raw bytes to ``str`` with best-effort encoding detection.

    Uses ``charset-normalizer`` first, then falls back to UTF-8 (replacing
    invalid bytes) so that malformed encodings never raise.

    Args:
        data: Raw bytes.

    Returns:
        The decoded text.
    """
    from charset_normalizer import from_bytes

    best = from_bytes(data).best()
    if best is not None:
        return str(best)
    return data.decode("utf-8", errors="replace")


def _extract_text_plain(data: bytes) -> str:
    """Extract text from a plain ``.txt`` file."""
    return _decode_text(data)


def _extract_text_markdown(data: bytes) -> str:
    """Extract human-readable text from a Markdown file.

    The Markdown is tokenised with ``markdown-it-py`` and the textual content of
    the tokens is concatenated, dropping most of the markup while preserving
    headings and paragraph text. If tokenisation fails for any reason, the raw
    decoded text is returned as a safe fallback.
    """
    raw = _decode_text(data)
    try:
        from markdown_it import MarkdownIt

        md = MarkdownIt("commonmark")
        tokens = md.parse(raw)
        parts: list[str] = []
        for token in tokens:
            if token.type == "inline" and token.content:
                parts.append(token.content)
            elif token.type in {"fence", "code_block"} and token.content:
                parts.append(token.content)
        text = "\n".join(parts).strip()
        return text or raw
    except Exception:  # pragma: no cover - defensive fallback
        logger.warning("Markdown tokenisation failed; using raw text.")
        return raw


def _extract_text_pdf(data: bytes, filename: str) -> str:
    """Extract text from a PDF using ``pypdf``.

    Args:
        data: Raw PDF bytes.
        filename: Source filename (for error messages).

    Returns:
        Concatenated text of all pages.

    Raises:
        ParseError: If the PDF cannot be read.
    """
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
    except (PdfReadError, OSError, ValueError) as exc:
        raise ParseError(f"Could not read PDF '{filename}': {exc}") from exc

    pages: list[str] = []
    for page_number, page in enumerate(reader.pages):
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # pragma: no cover - per-page resilience
            logger.warning("Failed to extract page %d of '%s': %s", page_number, filename, exc)
    return "\n\n".join(part for part in pages if part).strip()


def _extract_text_docx(data: bytes, filename: str) -> str:
    """Extract text from a ``.docx`` file using ``python-docx``.

    Both paragraph text and table cell text are collected so structured content
    is not lost.

    Args:
        data: Raw ``.docx`` bytes.
        filename: Source filename (for error messages).

    Returns:
        The extracted text.

    Raises:
        ParseError: If the document cannot be opened.
    """
    from docx import Document as DocxDocument
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = DocxDocument(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise ParseError(f"Could not read DOCX '{filename}': {exc}") from exc

    parts: list[str] = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                parts.append("\t".join(cells))
    return "\n".join(parts).strip()


def _extract_text_doc(data: bytes, filename: str) -> str:
    """Best-effort extraction of legacy ``.doc`` (Word 97-2003) files.

    There is no pure-Python parser for the binary ``.doc`` format that is both
    reliable and dependency-free, so this delegates to the ``antiword`` CLI if
    it is installed. When the converter is unavailable, a clear
    :class:`ParseError` is raised advising conversion to ``.docx``.

    Args:
        data: Raw ``.doc`` bytes.
        filename: Source filename (for error messages).

    Returns:
        The extracted text.

    Raises:
        ParseError: If no converter is available or conversion fails.
    """
    converter = shutil.which("antiword")
    if not converter:
        raise ParseError(
            f"Legacy '.doc' file '{filename}' requires the 'antiword' tool, which is not "
            "installed. Install antiword or convert the file to .docx/.pdf."
        )
    try:
        result = subprocess.run(
            [converter, "-"],
            input=data,
            capture_output=True,
            timeout=60,
            check=True,
        )
    except (subprocess.SubprocessError, OSError) as exc:
        raise ParseError(f"antiword failed to convert '{filename}': {exc}") from exc
    return result.stdout.decode("utf-8", errors="replace").strip()
